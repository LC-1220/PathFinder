import asyncio
import json
import re
from datetime import datetime

import pytest

import fastapi_app as app


class ReportRequest:
    def __init__(self, role, payload, method="POST"):
        self.session = {"user_id": "report-test-user", "role": role}
        self.payload = payload
        self.method = method
        self.query_params = payload if method == "GET" else {}

    async def json(self):
        return self.payload


class ReportCursor:
    def __init__(self, role):
        self.role = role
        self.query = ""

    def execute(self, query, params=()):
        self.query = query
        if "FROM users" in query and "LEFT JOIN LATERAL" in query:
            assert query.count("?") == len(params)
            assert query.count("{0,1}") == 2

    def fetchone(self):
        return self.role, True, False

    def fetchall(self):
        if "FROM report_card_uploads" in self.query:
            return [("extracted", 2, 1, datetime(2025, 3, 1))]
        if "FROM course_recommendation_history" in self.query:
            return [(7, "[]")]
        return [("STEM", 2, 2, 90.0, 2, 2, 0)]


class ReportConnection:
    def __init__(self, role):
        self.cursor_instance = ReportCursor(role)

    def cursor(self):
        return self.cursor_instance

    def close(self):
        pass


def _call_report(monkeypatch, role, payload):
    monkeypatch.setattr(app, "_db_conn", lambda: ReportConnection(role))
    response = asyncio.run(app.admin_generate_report(ReportRequest(role, payload)))
    return response.status_code, json.loads(response.body)


def test_report_generator_is_registered_with_fastapi():
    assert any(
        route.path == "/admin/reports/generate" and "POST" in route.methods
        for route in app.api_app.routes
    )
    assert any(
        route.path == "/admin/reports/download" and "POST" in route.methods
        for route in app.api_app.routes
    )


@pytest.mark.parametrize("role", [app.ROLE_ADMIN, app.ROLE_SEMI_ADMIN])
@pytest.mark.parametrize(
    ("report_type", "expected_title"),
    [
        ("student_summary", "Student Summary Report"),
        ("report_cards", "Report Card Status Report"),
        ("recommendations", "Recommendation Frequency Report"),
    ],
)
def test_admin_roles_can_generate_each_report(monkeypatch, role, report_type, expected_title):
    status, response = _call_report(monkeypatch, role, {"report_type": report_type})

    assert status == 200
    assert response["success"] is True
    assert response["report"]["title"] == expected_title
    assert isinstance(response["report"]["rows"], list)


@pytest.mark.parametrize("role", [app.ROLE_ADMIN, app.ROLE_SEMI_ADMIN])
def test_admin_roles_can_generate_all_reports(monkeypatch, role):
    status, result = _call_report(monkeypatch, role, {"report_type": "all_reports", "strand": "HT"})

    assert status == 200
    report = result["report"]
    assert report["title"] == "All Reports"
    assert [section["type"] for section in report["sections"]] == ["student_summary", "report_cards", "recommendations"]
    assert all(section["filters"]["strand"] == "HT" for section in report["sections"])


def test_report_generator_rejects_non_admin_roles(monkeypatch):
    status, response = _call_report(monkeypatch, "student", {"report_type": "student_summary"})

    assert status == 403
    assert response["success"] is False


@pytest.mark.parametrize(
    "payload",
    [
        {"report_type": "unknown"},
        {"report_type": "student_summary", "strand": "not-configured"},
        {"report_type": "report_cards", "from_date": "not-a-date"},
        {"report_type": "recommendations", "from_date": "2025-03-02", "to_date": "2025-03-01"},
    ],
)
def test_report_generator_rejects_invalid_filters(monkeypatch, payload):
    status, response = _call_report(monkeypatch, app.ROLE_SEMI_ADMIN, payload)

    assert status == 400
    assert response["success"] is False


@pytest.mark.parametrize("role", [app.ROLE_ADMIN, app.ROLE_SEMI_ADMIN])
@pytest.mark.parametrize("report_type", ["student_summary", "report_cards", "recommendations", "all_reports"])
def test_admin_roles_can_download_pdf(monkeypatch, role, report_type):
    monkeypatch.setattr(app, "_db_conn", lambda: ReportConnection(role))

    response = asyncio.run(app.admin_download_report(ReportRequest(role, {"report_type": report_type})))

    assert response.status_code == 200
    assert response.media_type == "application/pdf"
    assert response.body.startswith(b"%PDF-")
    assert response.body.rstrip().endswith(b"%%EOF")
    assert f"pathfinder-{report_type}-" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    if report_type == "all_reports":
        assert len(re.findall(rb"/Type\s*/Page\b", response.body)) >= 3


@pytest.mark.parametrize("role, payload, expected_status", [
    ("student", {"report_type": "student_summary"}, 403),
    (app.ROLE_SEMI_ADMIN, {"report_type": "invalid"}, 400),
])
def test_pdf_download_preserves_report_access_and_validation(monkeypatch, role, payload, expected_status):
    monkeypatch.setattr(app, "_db_conn", lambda: ReportConnection(role))

    response = asyncio.run(app.admin_download_report(ReportRequest(role, payload)))

    assert response.status_code == expected_status
    assert response.media_type == "application/json"


def test_long_pdf_report_spans_multiple_pages(monkeypatch):
    _, result = _call_report(monkeypatch, app.ROLE_ADMIN, {"report_type": "recommendations"})
    report = result["report"]
    report["rows"] = [
        {"course": f"Course {index} & Studies", "recommendations": index, "share": "1%"}
        for index in range(100)
    ]

    pdf = app._admin_report_pdf(report)

    assert len(re.findall(rb"/Type\s*/Page\b", pdf)) > 1


class ActivityCursor:
    rows = [
        (3, "Admin One", "admin@example.com", "admin_account_updated", "=HYPERLINK(\"x\")", "<b>details</b>", "10.0.0.1", datetime(2026, 10, 23, 2, 0), "admin"),
        (2, "Student Two", "student@example.com", "login", "", "", "", datetime(2026, 10, 22, 2, 0), "student"),
        (1, "Admin Three", "", "logout", "", "", "", None, "admin"),
    ]

    def __init__(self, role):
        self.role = role

    def execute(self, query, params=()):
        assert "FROM admin_activity_logs" in query or "FROM users" in query

    def fetchone(self):
        return self.role, True, False

    def fetchall(self):
        return self.rows


class ActivityConnection:
    def __init__(self, role):
        self.role = role

    def cursor(self):
        return ActivityCursor(self.role)

    def close(self):
        pass


def _export_activity(monkeypatch, payload, role=app.ROLE_SUPER_ADMIN):
    monkeypatch.setattr(app, "_db_conn", lambda: ActivityConnection(role))
    return asyncio.run(app.admin_export_activity(ReportRequest(role, payload)))


def test_activity_export_is_registered_with_fastapi():
    assert any(
        route.path == "/admin/activity/export" and "POST" in route.methods
        for route in app.api_app.routes
    )


def test_activity_export_pdf_contains_only_selected_entries(monkeypatch):
    response = _export_activity(monkeypatch, {"format": "pdf", "ids": [3, 1], "filter_summary": "Admin activity"})

    assert response.status_code == 200
    assert response.media_type == "application/pdf"
    assert response.body.startswith(b"%PDF-")
    assert 'filename="pathfinder-activity-log-' in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"


def test_activity_export_xlsx_keeps_order_labels_and_inert_text(monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook

    response = _export_activity(monkeypatch, {"format": "xlsx", "ids": [2, 3], "filter_summary": "Search: \"x\""})

    assert response.status_code == 200
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sheet = load_workbook(BytesIO(response.body)).active
    data = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=5)]
    assert [row[1] for row in data] == ["Admin One", "Student Two"]
    assert data[0][0] == "2026-10-23 10:00:00 AM"
    assert data[0][4] == "Admin account updated"
    assert data[1][4] == "Signed in"
    assert data[0][5] == "=HYPERLINK(\"x\")"
    assert sheet.cell(row=5, column=6).data_type == "s"
    assert 'Search: "x"' in sheet["A2"].value


def test_activity_export_is_full_admin_only(monkeypatch):
    response = _export_activity(monkeypatch, {"format": "pdf", "ids": [1]}, role=app.ROLE_SEMI_ADMIN)

    assert response.status_code == 403


@pytest.mark.parametrize("payload", [
    {"format": "csv", "ids": [1]},
    {"format": "pdf", "ids": "1"},
    {"format": "pdf", "ids": ["1"]},
    {"format": "pdf", "ids": [True]},
    {"format": "xlsx", "ids": list(range(app.ADMIN_ACTIVITY_LIMIT + 1))},
    [],
])
def test_activity_export_rejects_invalid_requests(monkeypatch, payload):
    response = _export_activity(monkeypatch, payload)

    assert response.status_code == 400
